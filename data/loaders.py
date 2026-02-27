# pharmacy_deserts/data/loaders.py
"""
Data loading functions for pharmacy desert analysis.
Consolidates all read_* functions for CSV/Excel/API data sources.
"""
import os
import pandas as pd
import numpy as np
import openpyxl
from pathlib import Path
from utils.cache import cache_data


# =============================================================================
# Financial / Census Data
# =============================================================================

@cache_data
def read_financial_data(file_path, skip_rows=0):
    """Read financial/income data from Census CSV."""
    df = pd.read_csv(file_path, skiprows=skip_rows)
    df = df[['NAME', 'S1901_C01_012E']]
    df['zip'] = df['NAME'].str.extract(r'(\d{5})')
    return df


@cache_data
def read_education_data_acs(year=2023, api_key=None):
    """
    Read ACS S1501 Educational Attainment by ZCTA.
    Returns % HS or lower (25+) and % less than HS (25+).
    """
    import requests
    base = f"https://api.census.gov/data/{year}/acs/acs5/subject"
    vars_ = ["NAME", "S1501_C02_007E", "S1501_C02_008E", "S1501_C02_009E", "S1501_C02_014E", "S1501_C02_015E"]
    params = {"get": ",".join(vars_), "for": "zip code tabulation area:*"}
    if api_key:
        params["key"] = api_key
    r = requests.get(base, params=params, timeout=120)
    r.raise_for_status()
    data = r.json()
    df = pd.DataFrame(data[1:], columns=data[0]).rename(columns={
        "zip code tabulation area": "zip",
        "S1501_C02_007E": "pct_less_9",
        "S1501_C02_008E": "pct_9to12_no_diploma",
        "S1501_C02_009E": "pct_hs_grad",
        "S1501_C02_014E": "pct_hs_or_higher",
        "S1501_C02_015E": "pct_ba_or_higher",
    })
    for c in ["pct_less_9", "pct_9to12_no_diploma", "pct_hs_grad", "pct_hs_or_higher", "pct_ba_or_higher"]:
        df[c] = pd.to_numeric(df[c], errors='coerce')
    df["zip"] = df["zip"].astype(str).str.zfill(5)
    df["edu_hs_or_lower_pct"] = df[["pct_less_9", "pct_9to12_no_diploma", "pct_hs_grad"]].sum(axis=1)
    df["edu_less_than_hs_pct"] = 100 - df["pct_hs_or_higher"]
    return df[["zip", "edu_hs_or_lower_pct", "edu_less_than_hs_pct", "pct_ba_or_higher"]]


# =============================================================================
# Health Data
# =============================================================================

@cache_data
def read_health_data(file_path, skip_rows=0):
    """Read health burden data from PLACES."""
    df = pd.read_csv(file_path, skiprows=skip_rows)
    df = df[['ZCTA5', 'GHLTH_CrudePrev']]
    df['ZCTA5'] = df['ZCTA5'].astype(str).str.split('.').str[0].str.zfill(5)
    return df


@cache_data
def read_hhi_excel(file_path, skip_rows=0):
    """
    Read Heat-Health Index (HHI) Excel data.
    Returns: DataFrame with zip, heat_hhb (from HHB_SCORE), nbe_score, hhi_overall
    """
    df = pd.read_excel(file_path, dtype={'ZCTA': str}, skiprows=skip_rows)
    if 'ZCTA' not in df.columns:
        raise ValueError("HHI Excel must contain 'ZCTA' column.")
    df['zip'] = (
        df['ZCTA'].astype(str).str.extract(r'(\d{5})')[0].fillna('').str.zfill(5)
    )
    out = pd.DataFrame({'zip': df['zip']})
    if 'HHB_SCORE' in df.columns:
        out['heat_hhb'] = pd.to_numeric(df['HHB_SCORE'], errors='coerce')
    if 'NBE_SCORE' in df.columns:
        out['nbe_score'] = pd.to_numeric(df['NBE_SCORE'], errors='coerce')
    if 'OVERALL_SCORE' in df.columns:
        out['hhi_overall'] = pd.to_numeric(df['OVERALL_SCORE'], errors='coerce')
    return out.dropna(subset=['zip']).drop_duplicates(subset=['zip'])


# =============================================================================
# Population Data
# =============================================================================

@cache_data
def read_population_data(file_path, skip_rows=10):
    """Read population density and location data."""
    df = pd.read_csv(file_path, skiprows=skip_rows)
    df.columns = [str(c).strip() for c in df.columns]
    lower = {c.lower(): c for c in df.columns}
    required = ["zip", "population", "density", "lat", "long"]
    if any(k not in lower for k in required):
        raise ValueError(f"Expected columns Zip, population, density, lat, long; got {df.columns[:10].tolist()}")
    zip_col = lower["zip"]
    pop_col = lower["population"]
    dens_col = lower["density"]
    lat_col = lower["lat"]
    lon_col = lower["long"]

    out = pd.DataFrame({
        "zip": df[zip_col].astype(str).str.extract(r"(\d{5})")[0].str.zfill(5),
        "population": df[pop_col].astype(str).str.replace(",", "", regex=False),
        "pop_density": df[dens_col].astype(str).str.replace(",", "", regex=False),
        "lat": df[lat_col],
        "lon": df[lon_col],
    })
    out["population"] = pd.to_numeric(out["population"], errors="coerce")
    out["pop_density"] = pd.to_numeric(out["pop_density"], errors="coerce")
    out["lat"] = pd.to_numeric(out["lat"], errors="coerce")
    out["lon"] = pd.to_numeric(out["lon"], errors="coerce")
    return (out.dropna(subset=["zip"])
              .groupby("zip", as_index=False)
              .agg({"population": "sum", "pop_density": "max", "lat": "first", "lon": "first"}))


@cache_data
def read_population_labels(file_path):
    """Read city/state labels for ZIPs."""
    df = pd.read_csv(file_path, skiprows=10)
    df.columns = [str(c).strip() for c in df.columns]
    lower = {c.lower(): c for c in df.columns}
    zip_col = lower.get("zip")
    city_col = lower.get("city")
    st_col = lower.get("st") or lower.get("state")
    if not zip_col:
        return pd.DataFrame(columns=["zip", "city", "state"])
    out = pd.DataFrame({"zip": df[zip_col].astype(str).str.extract(r"(\d{5})")[0].str.zfill(5)})
    if city_col:
        out["city"] = df[city_col].astype(str).str.strip()
    if st_col:
        out["state"] = df[st_col].astype(str).str.strip()
    return out.dropna(subset=["zip"]).drop_duplicates(subset=["zip"])


# =============================================================================
# Pharmacy Data
# =============================================================================

@cache_data
def read_pharmacy_data(file_path):
    """
    Read pharmacy data from either a CSV file or a directory of Excel files.
    If file_path is a directory, loads all subset*_Table1_filter.xlsm files.
    """
    path = Path(file_path)

    # Check if it's a directory first (before trying to read as CSV)
    if path.exists() and path.is_dir():
        print(f"Loading pharmacy data from directory: {path}")
        subset_files = sorted(path.glob("subset*_Table1_filter.xlsm"))

        if not subset_files:
            print(f"Warning: No subset*_Table1_filter.xlsm files found in {path}")
            return pd.DataFrame(columns=['zip', 'pharmacy_name', 'state'])

        all_data = []
        total_records = 0

        for file in subset_files:
            try:
                wb = openpyxl.load_workbook(file, read_only=True, data_only=True)
                if not wb.sheetnames:
                    print(f"  ✗ Skipped {file.name}: No sheets found")
                    continue

                sheet = wb[wb.sheetnames[0]]
                data = []
                for row in sheet.iter_rows(values_only=True):
                    data.append(row)

                if len(data) < 2:  # Need at least header + 1 row
                    print(f"  ✗ Skipped {file.name}: No data rows")
                    continue

                # Create DataFrame (first row is header)
                df_subset = pd.DataFrame(data[1:], columns=data[0])

                # Extract relevant columns
                cols_needed = {
                    'Provider Organization Name (Legal Business Name)': 'pharmacy_name',
                    'Provider Business Practice Location Address State Name': 'state',
                    'Short_ZIP': 'zip'
                }

                df_extracted = pd.DataFrame()
                for old_col, new_col in cols_needed.items():
                    if old_col in df_subset.columns:
                        df_extracted[new_col] = df_subset[old_col]
                    else:
                        print(f"  ✗ Warning: Column '{old_col}' not found in {file.name}")
                        df_extracted[new_col] = None

                # Clean ZIP codes
                df_extracted['zip'] = df_extracted['zip'].astype(str).str.extract(r'(\d{5})')[0].str.zfill(5)
                df_extracted = df_extracted.dropna(subset=['zip', 'pharmacy_name'])

                all_data.append(df_extracted)
                total_records += len(df_extracted)
                print(f"  ✓ Loaded {file.name}: {len(df_extracted):,} pharmacies")

            except Exception as e:
                print(f"  ✗ Error loading {file.name}: {e}")
                continue

        if not all_data:
            print("Warning: No valid pharmacy data found")
            return pd.DataFrame(columns=['zip', 'pharmacy_name', 'state'])

        combined = pd.concat(all_data, ignore_index=True)
        print(f"Total pharmacy records loaded: {len(combined):,}")
        print(f"Unique ZIPs with pharmacies: {combined['zip'].nunique()}")

        return combined

    # Otherwise, load from CSV (fallback)
    else:
        df = pd.read_csv(file_path)
        df.columns = [str(c).strip() for c in df.columns]
        zip_col = "ZIP" if "ZIP" in df.columns else "Zip" if "Zip" in df.columns else "zip"
        name_col = "NAME" if "NAME" in df.columns else "Name" if "Name" in df.columns else "name"
        state_col = "STATE" if "STATE" in df.columns else "State" if "State" in df.columns else "state"

        out = pd.DataFrame({
            "zip": df[zip_col].astype(str).str.extract(r"(\d{5})")[0].str.zfill(5),
            "pharmacy_name": df[name_col].astype(str),
            "state": df[state_col].astype(str) if state_col in df.columns else None
        })
        return out.dropna(subset=["zip"]).drop_duplicates()


# =============================================================================
# County Desert / HUD Crosswalk Data
# =============================================================================

@cache_data
def read_hud_zip_county_crosswalk(path, skip_rows=0):
    """
    Read HUD ZIP↔County crosswalk; return [zip, county, state, weight].
    Uses TOT_RATIO if present else RES_RATIO.
    """
    ext = os.path.splitext(path)[1].lower()
    df = (
        pd.read_excel(path, dtype=str, skiprows=skip_rows)
        if ext in (".xlsx", ".xls")
        else pd.read_csv(path, dtype=str, low_memory=False, skiprows=skip_rows)
    )
    cols = {c.lower(): c for c in df.columns}
    zip_col = cols.get("zip") or cols.get("zipcode") or cols.get("zip_code")
    county_col = cols.get("county") or cols.get("county_fips") or cols.get("fips")
    state_col = cols.get("state") or cols.get("stabbr") or cols.get("stusps")
    weight_col = next((cols[c] for c in ["tot_ratio", "total_ratio", "res_ratio"] if c in cols), None)
    if not (zip_col and county_col and weight_col):
        raise ValueError("Crosswalk must have ZIP, COUNTY, and TOT_RATIO/RES_RATIO.")
    out = pd.DataFrame({
        "zip": df[zip_col].astype(str).str.extract(r"(\d{5})")[0].str.zfill(5),
        "county": df[county_col].astype(str).str.extract(r"(\d{5})")[0].str.zfill(5),
        "state": (df[state_col] if state_col else pd.Series(index=df.index, dtype="object")),
        "weight": pd.to_numeric(df[weight_col], errors="coerce").fillna(0.0)
    })
    out = out.dropna(subset=["zip", "county"])
    return out[out["weight"] > 0]


@cache_data
def read_county_desert_csv(path, skip_rows=0):
    """
    Read county-level desert dataset; return [county, county_desert, drive_time_min, desert_pop_pct].
    """
    df = pd.read_csv(path, dtype=str, low_memory=False, skiprows=skip_rows)
    fips_col = next((c for c in df.columns if "fips" in c.lower()), None)
    if not fips_col:
        raise ValueError("County dataset must include a county FIPS column.")
    df["county"] = df[fips_col].astype(str).str.extract(r"(\d{5})")[0].str.zfill(5)

    flag_col = next((c for c in df.columns if c.lower() in ["desert", "is_desert", "desert_flag", "model1_pharm_desert", "pharm_desert"]), None)
    if flag_col:
        val = pd.to_numeric(df[flag_col], errors="coerce").fillna(0.0).clip(0, 1)
    else:
        score_col = next((c for c in df.columns if any(k in c.lower() for k in ["score", "index", "risk", "prob"])), None)
        if not score_col:
            raise ValueError("No desert flag/score column found in county dataset.")
        raw = pd.to_numeric(df[score_col], errors="coerce")
        val = (raw - raw.min()) / (raw.max() - raw.min()) if raw.max() > raw.min() else 0.0

    drive_time_col = next((c for c in df.columns if "drive_time" in c.lower() and "min" in c.lower()), None)
    pop_pct_col = next((c for c in df.columns if "desert_pop" in c.lower() and "pct" in c.lower()), None)
    drive_time = pd.to_numeric(df[drive_time_col], errors="coerce") if drive_time_col else pd.Series(index=df.index, dtype=float)
    desert_pop_pct = pd.to_numeric(df[pop_pct_col], errors="coerce") if pop_pct_col else pd.Series(index=df.index, dtype=float)

    # Make pop % consistently 0–100
    if desert_pop_pct.notna().any() and desert_pop_pct.max(skipna=True) <= 1.0:
        desert_pop_pct = desert_pop_pct * 100.0

    out = pd.DataFrame({
        "county": df["county"],
        "county_desert": val,
        "drive_time_min": drive_time,
        "desert_pop_pct": desert_pop_pct
    })
    return out.dropna(subset=["county"]).drop_duplicates(subset=["county"])


@cache_data
def downscale_county_to_zip(county_df, xwalk_df,
                            tiny_cutoff=0.01, min_coverage=0.60, threshold=0.50):
    """Downscale county-level desert metrics to ZIP-level using HUD crosswalk weights."""
    # filter tiny overlaps
    xw = xwalk_df[xwalk_df["weight"] >= tiny_cutoff].copy()
    totals = xw.groupby("zip", as_index=False)["weight"].sum().rename(columns={"weight": "zip_total"})

    m = xw.merge(county_df, on="county", how="left")
    m["has_desert"] = m["county_desert"].notna()
    m["has_drive"] = m["drive_time_min"].notna()
    m["has_pop"] = m["desert_pop_pct"].notna()

    m["w_desert"] = np.where(m["has_desert"], m["weight"] * m["county_desert"], 0.0)
    m["w_drive"] = np.where(m["has_drive"], m["weight"] * m["drive_time_min"], 0.0)
    m["w_pop"] = np.where(m["has_pop"], m["weight"] * m["desert_pop_pct"], 0.0)

    agg = (m.groupby("zip", as_index=False)
             .agg(zip_wsum=("weight", "sum"),
                  wmatch_desert=("weight", lambda s: s[m.loc[s.index, "has_desert"]].sum()),
                  wmatch_drive=("weight", lambda s: s[m.loc[s.index, "has_drive"]].sum()),
                  wmatch_pop=("weight", lambda s: s[m.loc[s.index, "has_pop"]].sum()),
                  wval_desert=("w_desert", "sum"),
                  wval_drive=("w_drive", "sum"),
                  wval_pop=("w_pop", "sum")))

    out = agg.merge(totals, on="zip", how="left")
    out["zip_total"] = out["zip_total"].replace(0, np.nan)

    out["cov_desert"] = (out["wmatch_desert"] / out["zip_total"]).clip(0, 1)
    out["cov_drive"] = (out["wmatch_drive"] / out["zip_total"]).clip(0, 1)
    out["cov_pop"] = (out["wmatch_pop"] / out["zip_total"]).clip(0, 1)

    renorm_desert = out["wmatch_desert"] > 0
    renorm_drive = out["wmatch_drive"] > 0
    renorm_pop = out["wmatch_pop"] > 0

    out.loc[renorm_desert, "zip_desert_share"] = out.loc[renorm_desert, "wval_desert"] / out.loc[renorm_desert, "wmatch_desert"]
    out.loc[renorm_drive, "zip_drive_time"] = out.loc[renorm_drive, "wval_drive"] / out.loc[renorm_drive, "wmatch_drive"]
    out.loc[renorm_pop, "zip_desert_pop_pct"] = out.loc[renorm_pop, "wval_pop"] / out.loc[renorm_pop, "wmatch_pop"]

    # dominant-county fallback
    dom = (m.sort_values(["zip", "weight"], ascending=[True, False])
             .drop_duplicates("zip")[["zip", "county_desert", "drive_time_min", "desert_pop_pct"]]
             .rename(columns={"county_desert": "dom_desert", "drive_time_min": "dom_drive", "desert_pop_pct": "dom_pop"}))
    out = out.merge(dom, on="zip", how="left")
    for col_out, col_dom in [("zip_desert_share", "dom_desert"), ("zip_drive_time", "dom_drive"), ("zip_desert_pop_pct", "dom_pop")]:
        need = out[col_out].isna() & out[col_dom].notna()
        out.loc[need, col_out] = out.loc[need, col_dom]

    # state median fallback
    county_state = (xw.groupby(["county", "state"], as_index=False)["weight"].sum()
                      .sort_values(["county", "weight"], ascending=[True, False])
                      .drop_duplicates("county")[["county", "state"]])
    cws = county_df.merge(county_state, on="county", how="left")
    state_med = (cws.dropna(subset=["state"]).groupby("state")
                   .agg(state_desert=("county_desert", "median"),
                        state_drive=("drive_time_min", "median"),
                        state_pop=("desert_pop_pct", "median")))

    zip_state = (xw.groupby(["zip", "state"], as_index=False)["weight"].sum()
                   .sort_values(["zip", "weight"], ascending=[True, False])
                   .drop_duplicates("zip")[["zip", "state"]])

    out = out.merge(zip_state, on="zip", how="left").merge(state_med, on="state", how="left")
    for col_out, col_state in [("zip_desert_share", "state_desert"), ("zip_drive_time", "state_drive"), ("zip_desert_pop_pct", "state_pop")]:
        need = out[col_out].isna() & out[col_state].notna()
        out.loc[need, col_out] = out.loc[need, col_state]

    out["zip_alloc_coverage"] = out["cov_desert"].fillna(0.0)
    out["zip_desert_flag"] = (out["zip_desert_share"] >= threshold).astype("Int64")

    return out[["zip", "zip_desert_share", "zip_desert_flag", "zip_drive_time", "zip_desert_pop_pct",
                "zip_alloc_coverage", "cov_drive", "cov_pop"]].rename(
        columns={"cov_drive": "zip_alloc_cov_drive", "cov_pop": "zip_alloc_cov_pop"}
    )


# =============================================================================
# Pharmacist Data
# =============================================================================

@cache_data
def load_all_pharmacist_data(data_dir="raw_data"):
    """
    Load and combine all subset pharmacist files (.xlsm format).
    Returns DataFrame with Short_ZIP, Combined, Award, Phone, and Address columns.
    """
    # Look for .xlsm files instead of .xlsx
    subset_files = sorted(Path(data_dir).glob("subset*_Table2_filter.xlsm"))

    if not subset_files:
        print("Warning: No subset pharmacist .xlsm files found")
        return pd.DataFrame(columns=['Short_ZIP', 'Combined', 'Award', 'Phone', 'Address'])

    print(f"Loading {len(subset_files)} pharmacist subset files...")

    all_data = []
    for file in subset_files:
        try:
            # Load the workbook
            xlsx = openpyxl.load_workbook(file, read_only=True, data_only=True)

            if not xlsx.sheetnames:
                print(f"  ✗ Skipped {file.name}: No sheets found")
                continue

            # Get first sheet data
            sheet = xlsx[xlsx.sheetnames[0]]
            data = list(sheet.values)

            if not data:
                print(f"  ✗ Skipped {file.name}: Empty sheet")
                continue

            # Create DataFrame with first row as headers
            df = pd.DataFrame(data[1:], columns=data[0])

            # Check for required columns (handle both "Short_ ZIP" and "Shorter ZIP")
            zip_col = None
            if 'Short_ ZIP' in df.columns:
                zip_col = 'Short_ ZIP'
            elif 'Shorter ZIP' in df.columns:
                zip_col = 'Shorter ZIP'
            elif 'Short_ZIP' in df.columns:
                zip_col = 'Short_ZIP'

            # Check for Award column
            award_col = 'Award?' if 'Award?' in df.columns else None

            # Check for Phone and Address columns
            phone_col = 'Provider Business Practice Location Address Telephone Number' if 'Provider Business Practice Location Address Telephone Number' in df.columns else None
            address_col = 'Provider First Line Business Practice Location Address' if 'Provider First Line Business Practice Location Address' in df.columns else None

            if zip_col and 'Combined' in df.columns:
                # Extract columns (include Award, Phone, Address if available)
                cols_to_extract = [zip_col, 'Combined']
                new_cols = ['Short_ZIP', 'Combined']

                if award_col:
                    cols_to_extract.append(award_col)
                    new_cols.append('Award')
                if phone_col:
                    cols_to_extract.append(phone_col)
                    new_cols.append('Phone')
                if address_col:
                    cols_to_extract.append(address_col)
                    new_cols.append('Address')

                df_subset = df[cols_to_extract].copy()
                df_subset.columns = new_cols

                # Add missing columns with None
                if 'Award' not in df_subset.columns:
                    df_subset['Award'] = None
                if 'Phone' not in df_subset.columns:
                    df_subset['Phone'] = None
                if 'Address' not in df_subset.columns:
                    df_subset['Address'] = None

                all_data.append(df_subset)

                extras = []
                if award_col:
                    extras.append("Award")
                if phone_col:
                    extras.append("Phone")
                if address_col:
                    extras.append("Address")
                extra_info = f" (with {', '.join(extras)})" if extras else ""
                print(f"  ✓ Loaded {file.name}: {len(df)} records{extra_info}")
            else:
                print(f"  ✗ Skipped {file.name}: Missing required columns")
                print(f"    Available: {df.columns.tolist()[:10]}")

        except Exception as e:
            print(f"  ✗ Error loading {file.name}: {e}")
            continue

    if not all_data:
        print("Warning: No valid pharmacist data found")
        return pd.DataFrame(columns=['Short_ZIP', 'Combined', 'Award', 'Phone', 'Address'])

    # Combine all subsets
    combined = pd.concat(all_data, ignore_index=True)

    # Remove any null values in required columns
    combined = combined.dropna(subset=['Short_ZIP', 'Combined'])

    # Count awarded pharmacists (Award=1 means they received an award)
    awarded_count = (combined['Award'] == 1).sum() if 'Award' in combined.columns else 0

    print(f"Total pharmacist records loaded: {len(combined)}")
    print(f"Unique ZIPs with pharmacists: {combined['Short_ZIP'].nunique()}")
    print(f"Pharmacists with awards: {awarded_count}")

    return combined


@cache_data
def get_pharmacists_for_zip(zip_code, pharmacist_df):
    """
    Get all pharmacists in a specific ZIP code.
    Returns list of tuples: (pharmacist_name, has_award, phone, address)
    """
    if pharmacist_df is None or pharmacist_df.empty:
        return []

    # Convert to string for matching
    zip_str = str(zip_code).strip()

    # Filter for this ZIP
    zip_data = pharmacist_df[pharmacist_df['Short_ZIP'].astype(str).str.strip() == zip_str]

    if zip_data.empty:
        return []

    # Build list of (name, has_award, phone, address) tuples
    pharmacists = []
    for _, row in zip_data.iterrows():
        name = row['Combined']
        if pd.notna(name) and str(name).strip():
            # Check if they have an award (Award column exists and value is 1)
            has_award = False
            if 'Award' in row.index and pd.notna(row['Award']):
                has_award = (int(row['Award']) == 1)

            # Get phone and address
            phone = str(row.get('Phone', '')).strip() if pd.notna(row.get('Phone')) else ''
            address = str(row.get('Address', '')).strip() if pd.notna(row.get('Address')) else ''

            # Format phone number nicely if it exists
            if phone and phone not in ['nan', 'None', '']:
                phone_digits = ''.join(filter(str.isdigit, phone))
                if len(phone_digits) == 10:
                    phone = f"({phone_digits[:3]}) {phone_digits[3:6]}-{phone_digits[6:]}"
                elif len(phone_digits) > 10:
                    phone = f"({phone_digits[-10:-7]}) {phone_digits[-7:-4]}-{phone_digits[-4:]}"
                else:
                    phone = phone_digits if phone_digits else ''
            else:
                phone = ''

            pharmacists.append((str(name).strip(), has_award, phone, address))

    # Sort: award winners first, then alphabetically by name
    pharmacists.sort(key=lambda x: (not x[1], x[0].lower()))

    return pharmacists
